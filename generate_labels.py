#!/usr/bin/env python3
"""Generate Avery 5160 labels (1\" x 2.625\", 3x10 per page) from responses-jumana-ben.xlsx.

One PDF per thaali day (only people who RSVP'd Yes for that day).
Label content: Name, City, Dish, "Date - Size".
Size taken from the "Thaali Size" column; Aamil Saheb gets Large for Thecha Chicken.
"""
import os
import sys
import openpyxl
from fpdf import FPDF

XLSX = "responses-jumana-ben.xlsx"
OUT_DIR = "labels"

# Avery 5160 layout (inches)
LABEL_W, LABEL_H = 2.625, 1.0
COLS, ROWS = 3, 10
MARGIN_L, MARGIN_T = 0.1875, 0.5     # left/top margin to first label
GAP_X = 0.125                         # horizontal gap between labels
# 5160 rows have ~0.0 gap; 3 cols: 2.625*3 + 0.125*2 + 0.1875*2 = 7.875+0.25+0.375=8.5 ok

# day -> (dish short, date label, filename slug)
DAYS = [
    ("Thaali day - Aug 4 - Smoked Chicken",                       "Smoked Chicken",   "Aug 4",  "aug04_smoked_chicken"),
    ("Thaali day - Aug 5 - Veg Korma",                            "Veg Korma",        "Aug 5",  "aug05_veg_korma"),
    ("Thaali day - Aug 6 - Thecha Chicken (Kharas, Chicken Drumsticks only)", "Thecha Chicken", "Aug 6",  "aug06_thecha_chicken"),
    ("Thaali day - Aug 7 - Chole Aaloo",                          "Chole Aaloo",      "Aug 7",  "aug07_chole_aaloo"),
    ("Thaali day - Aug 10 - Red Mutton Tarkari",                  "Red Mutton Tarkari","Aug 10", "aug10_red_mutton_tarkari"),
]

# Override: Aamil Saheb gets Large for Thecha Chicken (Aug 6), normal size otherwise.
THECHA_SLUG = "aug06_thecha_chicken"
OVERRIDE_LARGE = {"aamil saheb"}

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Sheet1"]
    headers = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    name_idx = col["Name2"]
    size_idx = col["Thaali Size"]
    city_idx = col["City of Residence"]

    total = 0
    for di, (hkey, dish, date, slug) in enumerate(DAYS):
        # people who RSVP'd Yes for this day: (name_key, name, city, size)
        day_labels = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[col[hkey]] != "Yes":
                continue
            name = (r[name_idx] or "").strip()
            if not name:
                continue
            city = (r[city_idx] or "").strip()
            sz = (r[size_idx] or "").strip()
            if slug == THECHA_SLUG and name.lower() in OVERRIDE_LARGE:
                sz = "Large"
            day_labels.append((name.lower(), name, city, sz))
        day_labels.sort(key=lambda x: x[0])

        pdf = FPDF(unit="in", format="Letter")
        pdf.set_auto_page_break(False)
        pdf.set_margins(0, 0)
        pdf.add_page()
        row_h = 1.0

        for i, (_, name, city, sz) in enumerate(day_labels):
            if i > 0 and i % (COLS * ROWS) == 0:
                pdf.add_page()
            idx = i % (COLS * ROWS)
            col_i = idx % COLS
            row_i = idx // COLS
            x = MARGIN_L + col_i * (LABEL_W + GAP_X)
            y = MARGIN_T + row_i * row_h
            pad = 0.08
            inner = LABEL_W - 2*pad
            # Name (bold)
            pdf.set_xy(x + pad, y + 0.08)
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(inner, 0.16, name, align="C")
            # City
            cy = pdf.get_y() + 0.02
            pdf.set_xy(x + pad, cy)
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(inner, 0.13, city, align="C")
            # Dish
            cy = pdf.get_y() + 0.04
            pdf.set_xy(x + pad, cy)
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(inner, 0.16, dish, align="C")
            # Date - Size (one line)
            cy = pdf.get_y() + 0.05
            pdf.set_xy(x + pad, cy)
            pdf.set_font("Helvetica", "B", 10)
            pdf.multi_cell(inner, 0.14, f"{date} - {sz}", align="C")

        out = os.path.join(OUT_DIR, f"{slug}.pdf")
        pdf.output(out)
        print(f"{out}: {len(day_labels)} labels ({-(-len(day_labels)//30)} sheet(s))", file=sys.stderr)
        total += len(day_labels)

    print(f"Total labels: {total}", file=sys.stderr)

if __name__ == "__main__":
    main()