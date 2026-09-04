#!/usr/bin/env python3
"""Build delivery coordinator workbook from responses-jumana-ben.xlsx.

- Master sheet mirroring the source.
- One sheet per thaali day: names grouped by size, with a per-city size
  count summary table (for quick driver counting) and a roti pickup table.
"""
import openpyxl, re
from collections import Counter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "responses-jumana-ben.xlsx"
OUT = "delivery_coordinator.xlsx"

DAYS = [
    ("Thaali day - Aug 4 - Smoked Chicken",                       "Aug 4 - Smoked Chicken"),
    ("Thaali day - Aug 5 - Veg Korma",                            "Aug 5 - Veg Korma"),
    ("Thaali day - Aug 6 - Thecha Chicken (Kharas, Chicken Drumsticks only)", "Aug 6 - Thecha Chicken"),
    ("Thaali day - Aug 7 - Chole Aaloo",                          "Aug 7 - Chole Aaloo"),
    ("Thaali day - Aug 10 - Red Mutton Tarkari",                  "Aug 10 - Red Mutton Tarkari"),
]
THECHA_KEY = DAYS[2][0]
OVERRIDE_LARGE = {"aamil saheb"}
CITIES = ["Abbotsford", "Burnaby", "Nanaimo", "Pitt Meadows", "Richmond", "Surrey", "Vancouver"]

def norm_city(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().title()

# styles
HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="305496")
SIZE_FONT = Font(bold=True, size=12, color="305496")
SIZE_FILL = PatternFill("solid", fgColor="D9E1F2")
SUB_FONT = Font(bold=True, size=11, color="FFFFFF")
SUB_FILL = PatternFill("solid", fgColor="548235")
ROTI_FONT = Font(bold=True, size=11, color="FFFFFF")
ROTI_FILL = PatternFill("solid", fgColor="BF8F00")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

def style_header(ws, row, ncols, font=HDR_FONT, fill=HDR_FILL):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = font
        cell.fill = fill
        cell.alignment = CENTER
        cell.border = BORDER

def width(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def main():
    src = openpyxl.load_workbook(SRC, data_only=True)
    sws = src["Sheet1"]
    hdr = [c.value for c in sws[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    name_idx = ci["Name2"]
    phone_idx = ci["Phone number"]
    city_idx = ci["City of Residence"]
    size_idx = ci["Thaali Size"]

    people = []
    for r in sws.iter_rows(min_row=2, values_only=True):
        name = (r[name_idx] or "").strip()
        if not name: continue
        people.append({
            "name": name,
            "phone": (r[phone_idx] or "").strip() if r[phone_idx] else "",
            "city": norm_city(r[city_idx]),
            "size": (r[size_idx] or "").strip() if r[size_idx] else "",
            "days": {hkey: (r[ci[hkey]] == "Yes") for hkey, _ in DAYS},
        })

    out = openpyxl.Workbook()

    # ---- Master sheet ----
    ms = out.active
    ms.title = "Master"
    ms.append(hdr)
    style_header(ms, 1, len(hdr))
    for r in sws.iter_rows(min_row=2, values_only=True):
        row = list(r)
        if row[city_idx] is not None:
            row[city_idx] = norm_city(row[city_idx])
        ms.append(row)
    for i, h in enumerate(hdr, 1):
        w = 14
        if h and len(str(h)) > 18: w = 28
        ms.column_dimensions[get_column_letter(i)].width = w
    ms.freeze_panes = "A2"

    # ---- per-day sheets ----
    for hkey, label in DAYS:
        ws = out.create_sheet(title=label[:31])
        # Title row
        ws.merge_cells("A1:F1")
        t = ws.cell(row=1, column=1, value=label)
        t.font = Font(bold=True, size=14, color="305496")
        t.alignment = CENTER

        # ============ LEFT: name list (cols A-D) ============
        headers = ["Size", "Name", "Phone", "City"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=2, column=i, value=h)
        style_header(ws, 2, 4)

        day_people = [p for p in people if p["days"][hkey]]
        if hkey == THECHA_KEY:
            day_people = [dict(p, size="Large") if p["name"].lower() in OVERRIDE_LARGE else p for p in day_people]
        order = {"Large": 0, "Medium": 1, "Small": 2}
        day_people.sort(key=lambda p: (order.get(p["size"], 9), p["city"], p["name"].lower()))

        row = 3
        cur_size = None
        for p in day_people:
            sz = p["size"] or "?"
            if sz != cur_size:
                cur_size = sz
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                sc = ws.cell(row=row, column=1, value=f"{sz}  ({sum(1 for x in day_people if (x['size'] or '?')==sz)} thaalis)")
                sc.font = SIZE_FONT
                sc.fill = SIZE_FILL
                sc.alignment = LEFT
                row += 1
            ws.append([sz, p["name"], p["phone"], p["city"]])
            for c in range(1, 5):
                ws.cell(row=row, column=c).border = BORDER
                ws.cell(row=row, column=c).alignment = LEFT if c in (2,3,4) else CENTER
            row += 1
        ws.append([])
        ws.append(["", f"Total: {len(day_people)} thaalis", "", ""])
        ws.cell(row=row+1, column=2).font = Font(bold=True)
        name_list_end = row + 1

        # ============ RIGHT: per-city size count (cols F-J) ============
        sc_col = 6
        ws.merge_cells(start_row=2, start_column=sc_col, end_row=2, end_column=sc_col+4)
        sh = ws.cell(row=2, column=sc_col, value="Per-City Size Count")
        sh.font = Font(bold=True, color="FFFFFF")
        sh.fill = PatternFill("solid", fgColor="548235")
        sh.alignment = CENTER

        sub_headers = ["City", "Large", "Medium", "Small", "Total"]
        for i, h in enumerate(sub_headers):
            ws.cell(row=3, column=sc_col+i, value=h)
        style_header(ws, 3, sc_col+len(sub_headers)-1)

        # compute counts
        city_counts = {c: Counter() for c in CITIES}
        for p in day_people:
            city_counts.setdefault(p["city"], Counter())[p["size"] or "?"] += 1

        cr = 4
        for city in CITIES:
            cc = city_counts.get(city, Counter())
            ws.append([None]*5)  # keep row aligned
            ws.cell(row=cr, column=sc_col, value=city)
            ws.cell(row=cr, column=sc_col+1, value=cc.get("Large", 0))
            ws.cell(row=cr, column=sc_col+2, value=cc.get("Medium", 0))
            ws.cell(row=cr, column=sc_col+3, value=cc.get("Small", 0))
            ws.cell(row=cr, column=sc_col+4, value=cc.get("Large",0)+cc.get("Medium",0)+cc.get("Small",0))
            for c in range(sc_col, sc_col+5):
                ws.cell(row=cr, column=c).border = BORDER
                ws.cell(row=cr, column=c).alignment = LEFT if c == sc_col else CENTER
            cr += 1
        # totals row
        ws.cell(row=cr, column=sc_col, value="TOTAL")
        for i, key in enumerate(["Large","Medium","Small"]):
            ws.cell(row=cr, column=sc_col+1+i, value=sum(city_counts[c].get(key,0) for c in CITIES))
        ws.cell(row=cr, column=sc_col+4, value=len(day_people))
        for c in range(sc_col, sc_col+5):
            ws.cell(row=cr, column=c).font = Font(bold=True)
            ws.cell(row=cr, column=c).fill = PatternFill("solid", fgColor="D9E1F2")
            ws.cell(row=cr, column=c).border = BORDER
            ws.cell(row=cr, column=c).alignment = LEFT if c == sc_col else CENTER

        # ============ Roti pickup table (below city count) ============
        cr += 3
        ws.merge_cells(start_row=cr, start_column=sc_col, end_row=cr, end_column=sc_col+3)
        rh = ws.cell(row=cr, column=sc_col, value="Roti Pickup")
        rh.font = ROTI_FONT
        rh.fill = ROTI_FILL
        rh.alignment = CENTER
        cr += 1
        roti_headers = ["City", "Roti Caterer", "Pickup Location", "Phone"]
        for i, h in enumerate(roti_headers):
            ws.cell(row=cr, column=sc_col+i, value=h)
        style_header(ws, cr, sc_col+len(roti_headers)-1)
        cr += 1
        for city in CITIES:
            ws.cell(row=cr, column=sc_col, value=city)
            for c in range(sc_col, sc_col+4):
                ws.cell(row=cr, column=c).border = BORDER
                ws.cell(row=cr, column=c).alignment = LEFT if c == sc_col else CENTER
            cr += 1

        width(ws, [10, 28, 16, 16, 3, 16, 10, 10, 10, 12, 14, 22, 16])
        ws.freeze_panes = "A3"

    out.save(OUT)
    print(f"Wrote {OUT}")

if __name__ == "__main__":
    main()