#!/usr/bin/env python3
"""Generate Avery 5160 labels with player names and jersey sizes from Intercity Sports Festival Size Charts.xlsx.

One label per row (rows are printed as-is). Label content: Name, City, "Size: X".
"""
import os
import sys
import openpyxl
from fpdf import FPDF

XLSX = "Intercity Sports Festival Size Charts.xlsx"
OUT = "jersey_labels.pdf"

# Avery 5160 layout (inches)
LABEL_W, LABEL_H = 2.625, 1.0
COLS, ROWS = 3, 10
MARGIN_L, MARGIN_T = 0.1875, 0.5
GAP_X = 0.125


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Sheet1"]
    headers = [c.value for c in ws[1]]
    name_idx = headers.index("Name")
    city_idx = headers.index("City ")
    size_idx = headers.index("Size")

    players = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = (r[name_idx] or "").strip()
        if not name:
            continue
        city = (r[city_idx] or "").strip() if r[city_idx] else ""
        sz = (r[size_idx] or "").strip() if r[size_idx] else ""
        players.append((name, city, sz))

    pdf = FPDF(unit="in", format="Letter")
    pdf.set_auto_page_break(False)
    pdf.set_margins(0, 0)
    pdf.add_page()

    for i, (name, city, sz) in enumerate(players):
        if i > 0 and i % (COLS * ROWS) == 0:
            pdf.add_page()
        idx = i % (COLS * ROWS)
        col_i = idx % COLS
        row_i = idx // COLS
        x = MARGIN_L + col_i * (LABEL_W + GAP_X)
        y = MARGIN_T + row_i * LABEL_H
        pad = 0.08
        inner = LABEL_W - 2 * pad
        # Name (bold)
        pdf.set_xy(x + pad, y + 0.08)
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(inner, 0.16, name, align="C")
        # City
        cy = pdf.get_y() + 0.02
        pdf.set_xy(x + pad, cy)
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(inner, 0.13, city, align="C")
        # Size (bold)
        cy = pdf.get_y() + 0.06
        pdf.set_xy(x + pad, cy)
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(inner, 0.18, f"Size: {sz}", align="C")

    pdf.output(OUT)
    print(f"{OUT}: {len(players)} labels ({-(-len(players)//30)} sheet(s))", file=sys.stderr)


if __name__ == "__main__":
    main()
